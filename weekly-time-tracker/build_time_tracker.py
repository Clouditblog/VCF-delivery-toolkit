"""Builds weekly_time_tracker.xlsx: project rows, Mon-Fri hours, 8h/day and 40h/week targets."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule

OUT = sys.argv[1] if len(sys.argv) > 1 else "weekly_time_tracker.xlsx"

wb = Workbook()
ws = wb.active
ws.title = "Weekly Time"

F = "Arial"
hdr_fill = PatternFill("solid", fgColor="1F3864")
tot_fill = PatternFill("solid", fgColor="D9E1F2")
input_fill = PatternFill("solid", fgColor="FFFFE0")
thin = Side(style="thin", color="A6A6A6")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

ws["A1"] = "Weekly Project Time Tracker"
ws["A1"].font = Font(name=F, size=14, bold=True)
ws["A2"] = "Week commencing:"
ws["A2"].font = Font(name=F, bold=True)
ws["B2"].fill = input_fill
ws["B2"].border = border
ws["B2"].font = Font(name=F, color="0000FF")
ws["B2"].number_format = "dd mmm yyyy"

headers = ["Project", "Allocated (hrs)", "Remaining (hrs)", "Mon", "Tue", "Wed", "Thu", "Fri", "Total (hrs)"]
HR = 4
for i, h in enumerate(headers, 1):
    c = ws.cell(row=HR, column=i, value=h)
    c.font = Font(name=F, bold=True, color="FFFFFF")
    c.fill = hdr_fill
    c.alignment = Alignment(horizontal="center" if i > 1 else "left", vertical="center")
    c.border = border

first, last = HR + 1, HR + 10   # 10 project rows
for r in range(first, last + 1):
    for col in range(1, 10):
        c = ws.cell(row=r, column=col)
        c.font = Font(name=F)
        c.border = border
        if col in (1, 2, 4, 5, 6, 7, 8):
            c.fill = input_fill
            c.font = Font(name=F, color="0000FF")
        if col > 1:
            c.number_format = "0.0;-0.0;-"
            c.alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=3, value=f"=B{r}-I{r}")
    ws.cell(row=r, column=9, value=f"=SUM(D{r}:H{r})")

# example row
name, alloc, days = "Example - VCF upgrade", 12, (3, 2, 2.5, 2, 1.5)
ws.cell(row=first, column=1, value=name)
ws.cell(row=first, column=2, value=alloc)
for i, d in enumerate(days):
    ws.cell(row=first, column=4 + i, value=d)

# totals row
tr = last + 1
ws.cell(row=tr, column=1, value="Total")
for col in range(1, 10):
    c = ws.cell(row=tr, column=col)
    c.font = Font(name=F, bold=True)
    c.fill = tot_fill
    c.border = border
    if col > 1:
        L = "ABCDEFGHI"[col - 1]
        c.value = f"=SUM({L}{first}:{L}{last})"
        c.number_format = "0.0;-0.0;0.0"
        c.alignment = Alignment(horizontal="center")

green = PatternFill("solid", bgColor="C6EFCE", fgColor="C6EFCE")
red = PatternFill("solid", bgColor="FFC7CE", fgColor="FFC7CE")
amber = PatternFill("solid", bgColor="FFEB9C", fgColor="FFEB9C")

# day totals: 8 = green, over 8 = red, under = amber
ws.conditional_formatting.add(f"D{tr}:H{tr}", CellIsRule(operator="equal", formula=["8"], fill=green))
ws.conditional_formatting.add(f"D{tr}:H{tr}", CellIsRule(operator="greaterThan", formula=["8"], fill=red))
ws.conditional_formatting.add(f"D{tr}:H{tr}", CellIsRule(operator="lessThan", formula=["8"], fill=amber))
# week total and allocated total: 40
for rng in (f"B{tr}", f"I{tr}"):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=["40"], fill=green))
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["40"], fill=red))
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["40"], fill=amber))
# negative remaining = over budget
ws.conditional_formatting.add(f"C{first}:C{last}", CellIsRule(operator="lessThan", formula=["0"], fill=red))

# legend
lr = tr + 2
notes = [
    "How to use",
    "Yellow cells with blue text are yours to fill in: project name, allocated hours for the week, and hours worked each day.",
    "Remaining = Allocated minus Total. It turns red if you go over budget.",
    "Total row: each day should reach 8 hours (green), and the Allocated and Total columns should each reach 40 hours (green).",
    "Amber means under target, red means over. Hours can be decimals, e.g. 1.5 for 90 minutes.",
    "The first row is an example. Overwrite or clear it. Add more project rows by inserting a row inside the table.",
]
for i, t in enumerate(notes):
    c = ws.cell(row=lr + i, column=1, value=t)
    c.font = Font(name=F, bold=(i == 0), size=10)

ws.column_dimensions["A"].width = 28
for L in "BCDEFGHI":
    ws.column_dimensions[L].width = 15 if L in "BCI" else 9
ws.row_dimensions[HR].height = 22
ws.freeze_panes = f"B{HR+1}"
ws.sheet_view.showGridLines = False

wb.save(OUT)
print("wrote", OUT)
